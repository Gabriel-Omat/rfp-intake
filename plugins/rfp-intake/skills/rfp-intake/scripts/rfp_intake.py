#!/usr/bin/env python3
"""Local RFP evidence store, review ledger, wiki publisher, and context retrieval.

The CLI extracts source data. An agent or human supplies semantic review batches.
No source files are edited; no model call or approval is synthesized.
"""
from __future__ import annotations

import argparse
import collections
import contextlib
import csv
import datetime as dt
import difflib
import hashlib
import html
import json
import os
from pathlib import Path
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import time
import unicodedata
from urllib.parse import quote

if os.name == "nt":
    import msvcrt
else:
    import fcntl

VERSION = "1.0.0"
SOURCE = "00_SOURCE OF RECORD"
INTAKE = "01_INTAKE"
EXTRACT = "02_FAITHFUL EXTRACTION"
KNOWLEDGE = "03_PROJECT KNOWLEDGE"
TYPES = {"scope": "SC", "requirement": "REQ", "date": "DATE", "bid_item": "BID",
         "form": "FORM", "amendment": "AMD", "risk": "RISK", "rfi": "RFI",
         "drawing": "DRAW", "submittal": "SUB", "deliverable": "DEL"}
FILES = {"scope": "scope_register", "requirement": "requirements_register",
         "date": "dates_and_deadlines", "bid_item": "bid_items", "form": "forms_checklist",
         "amendment": "amendment_log", "risk": "risks_and_conflicts", "rfi": "questions_and_RFIs",
         "drawing": "drawing_register", "submittal": "submittal_register", "deliverable": "deliverables_register"}
LAYER = {"idiq", "seed_task_order", "task_order", "project", "unknown"}
RELATIONS = {"applies_to", "prices", "requires", "conflicts_with", "clarifies",
             "candidate_supersedes", "supersedes", "supports", "depicts"}
ATTESTATIONS = {"package_complete", "authority_reconciled", "critical_records_reviewed", "release_to_estimating"}


def now():
    return dt.datetime.now(dt.timezone.utc).isoformat()


def sha(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()


def identifier(prefix, value):
    return prefix + "-" + digest(value)[:16]


def norm(value):
    return " ".join(unicodedata.normalize("NFKC", str(value)).split())


def safe_json(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def write_text(path, text):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", newline="", dir=path.parent, delete=False) as f:
        f.write(text)
        tmp = f.name
    replace_file(tmp, path)


def replace_file(tmp, path):
    """os.replace with retries: on Windows, antivirus, sync clients, or an open viewer can briefly lock a file."""
    for attempt in range(10):
        try:
            os.replace(tmp, path)
            return
        except PermissionError:
            if os.name != "nt" or attempt == 9:
                with contextlib.suppress(OSError):
                    os.remove(tmp)
                raise PermissionError(f"Cannot replace {path}; close any program that has it open and retry")
            time.sleep(0.2 * (attempt + 1))


def rel_path(path, base):
    """Stored relative paths always use forward slashes so packets are portable between Windows, macOS, and Linux."""
    return Path(path).relative_to(base).as_posix()


def dump(path, obj):
    write_text(path, json.dumps(obj, indent=2, ensure_ascii=False, default=safe_json) + "\n")


def jsonl(path, rows):
    write_text(path, "".join(json.dumps(x, ensure_ascii=False, default=safe_json) + "\n" for x in rows))


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def csv_write(path, rows, fields):
    import io
    out = io.StringIO(newline="")
    writer = csv.DictWriter(out, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        values = {}
        for k in fields:
            v = row.get(k)
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False)
            # Human-facing CSVs are formula-safe. JSON keeps the exact typed value.
            if isinstance(v, str) and v.startswith(("=", "+", "-", "@", "\t", "\r")):
                v = "'" + v
            values[k] = v
        writer.writerow(values)
    write_text(path, out.getvalue())


@contextlib.contextmanager
def project_lock(project):
    project.mkdir(parents=True, exist_ok=True)
    with (project / ".intake.lock").open("a+") as f:
        try:
            if os.name == "nt":
                f.seek(0)
                msvcrt.locking(f.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                fcntl.flock(f, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError as exc:
            raise OSError(f"Project packet is in use by another intake command: {project}") from exc
        try:
            yield
        finally:
            if os.name == "nt":
                f.seek(0)
                with contextlib.suppress(OSError):
                    msvcrt.locking(f.fileno(), msvcrt.LK_UNLCK, 1)


def load(project):
    return read_json(project / INTAKE / "state.json")


def save(project, state):
    state["updated_at"] = now()
    dump(project / INTAKE / "state.json", state)


def blank_state(name):
    return {"schema_version": VERSION, "project_name": name, "snapshot": None, "files": [],
            "documents": {}, "evidence": {}, "sections": {}, "records": {}, "relationships": {},
            "expected_documents": {}, "issues": {}, "section_reviews": {}, "attestations": {},
            "batches": {}, "ocr": {}, "estimate_links": {}, "change_reports": []}


def add_issue(state, kind, target, message, severity="blocking"):
    key = identifier("EXC", [kind, target, message])
    if key not in state["issues"]:
        state["issues"][key] = {"issue_id": key, "kind": kind, "target_id": target,
                                "message": message, "severity": severity, "created_at": now()}
    return key


def active_docs(state):
    return {f["document_id"] for f in state["files"] if f.get("document_id")}


def active_evidence(state):
    docs = active_docs(state)
    return {k: v for k, v in state["evidence"].items() if v["document_id"] in docs}


def classify_filename(name):
    n = name.lower().replace("+", " ").replace("_", " ")
    if "amend" in n or "addend" in n:
        return "amendment_candidate"
    if "solicitation" in n:
        return "solicitation"
    if "specification" in n:
        return "specifications"
    if "pricing" in n or "price" in n:
        return "pricing_schedule"
    if "question" in n:
        return "qa_template"
    if "sow" in n or "statement" in n:
        return "statement_of_work"
    if "plan" in n or "drawing" in n:
        return "drawing_candidate"
    return "unclassified"


def scan_source(source, project, excluded):
    excludes = [(source / e).resolve() for e in excluded]
    paths = []
    for base, dirs, names in os.walk(source, followlinks=False):
        kept = []
        for d in sorted(dirs):
            p = Path(base) / d
            if d in {".git", ".codex", ".agents", ".claude", "__pycache__", "node_modules"}:
                continue
            if p.resolve().is_relative_to(project) or any(p.resolve().is_relative_to(x) for x in excludes):
                continue
            if p.is_symlink():
                paths.append(p)
            else:
                kept.append(d)
        dirs[:] = kept
        for name in sorted(names):
            p = Path(base) / name
            if name == ".DS_Store" or name.startswith("~$"):
                continue
            if not any(p.resolve().is_relative_to(x) for x in excludes):
                paths.append(p)
    return sorted(paths)


def printed_ref(text):
    matches = re.findall(r"SECTION\s+(\d{2}\s+\d{2}\s+\d{2}(?:\.\d{2})?)\s+Page\s+(\d+)", text, re.I)
    if matches:
        sec, page = matches[-1]
        return norm(sec), page, "section"
    matches = re.findall(r"Page\s+(\d+)\s+of\s+(\d+)", text, re.I)
    return (None, matches[-1][0], "document_or_appendix") if matches else (None, None, "unlabeled")


def bookmark_map(reader):
    result = {}
    def walk(items, depth=0):
        for item in items:
            if isinstance(item, list):
                walk(item, depth + 1)
            else:
                try:
                    p = reader.get_destination_page_number(item) + 1
                    title = str(item.title)
                    if depth == 0 or title == "PROJECT TABLE OF CONTENTS":
                        result[p] = result[p] + " / " + title if p in result else title
                except Exception:
                    pass
    walk(reader.outline)
    return result


def refine_sections(state):
    """Index within-page SOW/solicitation headings; retain full shared pages as context."""
    for did, doc in state["documents"].items():
        if doc.get("role") not in {"solicitation", "statement_of_work"}:
            continue
        pages = sorted([e for e in state["evidence"].values() if e["document_id"] == did and e["method"] == "pdfplumber_native"],
                       key=lambda e: e["pdf_page"])
        groups = {}
        current = None
        for e in pages:
            headings = []
            is_toc = "TABLE OF CONTENTS" in e["text"].upper() or len(re.findall(r"(?:\.{2,}|…+)\s*\d+\s*$", e["text"], re.M)) > 3
            if is_toc:
                headings = [(0, "Table of contents")]
            for match in ([] if is_toc else re.finditer(r"(?m)^(Section\s+[A-M]\s*[-–].+|(?:\d+\.\d+|\d+\.)\s+[A-Z][^\n]{0,110}|APPENDIX\s+[A-Z]:?)$", e["text"])):
                title = match.group().strip()
                if re.search(r"(?:\.{2,}|…+)\s*\d+$", title):
                    continue
                if title.startswith("Section ") or title.startswith("APPENDIX"):
                    headings.append((match.start(), title))
                else:
                    body = re.sub(r"^(?:\d+\.\d+|\d+\.)\s+", "", title)
                    letters = [c for c in body if c.isalpha()]
                    if (letters and sum(c.isupper() for c in letters) / len(letters) > .7) or (len(body.split()) <= 7 and not body.endswith(".")):
                        headings.append((match.start(), title))
            assignments = []
            if current and (not headings or headings[0][0] > 100):
                assignments.append(current)
            for _, title in headings:
                sid = identifier("SEC", [doc["sha256"], "semantic_heading", title])
                groups.setdefault(sid, {"section_id":sid,"document_id":did,"heading":title,"evidence_ids":[],
                                        "method":"detected_heading_with_shared_page_context"})
                assignments.append(sid)
                current = sid
            if not assignments:
                if current is None:
                    current = identifier("SEC", [doc["sha256"], "semantic_front_matter"])
                    groups.setdefault(current, {"section_id":current,"document_id":did,"heading":"Front matter",
                                                 "evidence_ids":[],"method":"detected_heading_with_shared_page_context"})
                assignments.append(current)
            for sid in dict.fromkeys(assignments):
                groups[sid]["evidence_ids"].append(e["evidence_id"])
            e["section_id"] = assignments[-1]
            e["section_ids"] = list(dict.fromkeys(assignments))
            e["heading"] = " / ".join(groups[s]["heading"] for s in e["section_ids"])
        if not pages:
            continue
        for e in state["evidence"].values():
            if e.get("parent_evidence_id") in state["evidence"] and e["document_id"] == did:
                parent = state["evidence"][e["parent_evidence_id"]]
                e["section_id"] = parent["section_id"]
                for sid in parent.get("section_ids", [parent["section_id"]]):
                    groups[sid]["evidence_ids"].append(e["evidence_id"])
        state["sections"] = {sid:s for sid,s in state["sections"].items() if s["document_id"] != did}
        state["sections"].update(groups)


def extract_pdf(project, state, doc):
    from pypdf import PdfReader
    import pdfplumber
    path = project / doc["source_path"]
    reader = PdfReader(path)
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError("PDF requires a password; source was not decrypted")
    doc.update(page_count=len(reader.pages), encrypted=reader.is_encrypted,
               metadata={str(k): safe_json(v) for k, v in (reader.metadata or {}).items()})
    bookmarks = bookmark_map(reader)
    doc["bookmarks"] = bookmarks
    root = project / EXTRACT / doc["document_id"]
    rows, table_rows = [], []
    section_key, heading = "front_matter", "Front matter"
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            evid = identifier("EV", [doc["sha256"], "pdf", i])
            error = None
            try:
                text = page.extract_text(layout=False) or ""
            except Exception as exc:
                text, error = "", str(exc)
            section, label, scheme = printed_ref(text)
            if section:
                section_key, heading = section, "UFGS " + section
            elif i in bookmarks:
                section_key, heading = "bookmark-" + str(i), bookmarks[i]
            elif re.search(r"APPENDIX\s+[A-Z]", text[:250]):
                section_key, heading = "appendix-" + str(i), norm(text[:160])
            secid = identifier("SEC", [doc["sha256"], section_key])
            imgs = len(page.images)
            chars = len(text.strip())
            typ = "native_text"
            if chars == 0 and not imgs:
                typ = "blank_candidate"
            elif imgs and chars < 300:
                typ = "image_heavy"
            elif imgs:
                typ = "mixed"
            elif chars < 80:
                typ = "sparse_text"
            item = {"evidence_id": evid, "document_id": doc["document_id"], "source_sha256": doc["sha256"],
                    "pdf_page": i, "printed_page": label, "printed_page_scheme": scheme,
                    "section_number": section, "heading": heading, "section_id": secid,
                    "sheet_name": None, "cell_or_range": None, "text": text, "method": "pdfplumber_native",
                    "page_type": typ, "char_count": chars, "image_count": imgs,
                    "page_width": float(page.width), "page_height": float(page.height),
                    "text_quality": "unverified", "error": error,
                    "text_path": rel_path(root / "pages" / f"{i:04}.txt", project)}
            write_text(project / item["text_path"], text + "\n")
            item["text_file_sha256"] = sha(project / item["text_path"])
            rows.append(item)
            state["evidence"][evid] = item
            sec = state["sections"].setdefault(secid, {"section_id": secid, "document_id": doc["document_id"],
                                                       "heading": heading, "evidence_ids": [], "method": "structural_candidate"})
            sec["evidence_ids"].append(evid)
            if typ in {"blank_candidate", "image_heavy", "mixed", "sparse_text"}:
                add_issue(state, "visual_review", evid, f"PDF page {i}: {typ}; inspect rendering and record disposition.")
            if error:
                add_issue(state, "extraction_error", evid, error)
            # Layout words preserve location, even if row/table reconstruction fails.
            words = page.extract_words() if not error else []
            jsonl(root / "words" / f"{i:04}.jsonl", words)
            # Only attempt ruled-table detection where page geometry suggests a table.
            if len(page.lines) + len(page.rects) >= 4:
                try:
                    for ti, table in enumerate(page.find_tables(), 1):
                        cells = table.extract()
                        tid = identifier("TABLE", [evid, ti])
                        table_rows.append({"table_id": tid, "evidence_id": evid, "pdf_page": i,
                                           "bbox": list(table.bbox), "rows": cells, "review_status": "unverified"})
                        add_issue(state, "table_review", evid,
                                  f"Table {ti} on PDF page {i}: compare row/column alignment and unassigned text with rendering.")
                except Exception as exc:
                    add_issue(state, "table_extraction_error", evid, str(exc))
            if i % 50 == 0:
                print(f"Extracted {i}/{len(pdf.pages)} pages: {doc['original_filename']}", file=sys.stderr, flush=True)
    jsonl(root / "pages.jsonl", rows)
    jsonl(root / "tables.jsonl", table_rows)
    doc["table_count"] = len(table_rows)
    doc["extraction_status"] = "extracted_unverified"
    doc["classification_counts"] = dict(collections.Counter(x["page_type"] for x in rows))
    # Check section-local printed sequences without treating unrelated references as requirements.
    grouped = collections.defaultdict(list)
    for r in rows:
        if r["section_number"]:
            grouped[r["section_number"]].append(int(r["printed_page"]))
    doc["section_page_checks"] = []
    for sec, nums in grouped.items():
        missing = sorted(set(range(1, max(nums) + 1)) - set(nums))
        duplicate = [n for n, count in collections.Counter(nums).items() if count > 1]
        doc["section_page_checks"].append({"section": sec, "observed": len(nums), "last_printed": max(nums),
                                           "missing": missing, "duplicates": duplicate})
        if missing or duplicate:
            add_issue(state, "page_sequence", doc["document_id"], f"{sec}: missing={missing}; duplicate={duplicate}")


def extract_workbook(project, state, doc):
    import openpyxl
    path = project / doc["source_path"]
    wb = openpyxl.load_workbook(path, data_only=False, read_only=False, keep_links=True)
    cached = openpyxl.load_workbook(path, data_only=True, read_only=False, keep_links=True)
    root = project / EXTRACT / doc["document_id"]
    cells, rows, sheets = [], [], []
    for ws in wb.worksheets:
        secid = identifier("SEC", [doc["sha256"], "sheet", ws.title])
        state["sections"][secid] = {"section_id": secid, "document_id": doc["document_id"],
                                    "heading": ws.title, "evidence_ids": [], "method": "worksheet"}
        formulas = 0
        for row in ws.iter_rows():
            populated = []
            for c in row:
                if c.value is None:
                    continue
                v = safe_json(c.value)
                formula = v if c.data_type == "f" else None
                formulas += bool(formula)
                cells.append({"sheet": ws.title, "cell": c.coordinate, "raw_value": v, "data_type": c.data_type,
                              "formula": formula, "cached_value": safe_json(cached[ws.title][c.coordinate].value),
                              "number_format": c.number_format, "comment": c.comment.text if c.comment else None,
                              "hyperlink": c.hyperlink.target if c.hyperlink else None})
                populated.append((c.coordinate, v))
            if not populated:
                continue
            rn = row[0].row
            eid = identifier("EV", [doc["sha256"], "xlsx", ws.title, rn])
            text = "\n".join(f"{coord}: {val}" for coord, val in populated)
            item = {"evidence_id": eid, "document_id": doc["document_id"], "source_sha256": doc["sha256"],
                    "pdf_page": None, "printed_page": None, "printed_page_scheme": None,
                    "section_number": None, "heading": ws.title, "section_id": secid,
                    "sheet_name": ws.title, "cell_or_range": f"A{rn}:{openpyxl.utils.get_column_letter(ws.max_column)}{rn}",
                    "text": text, "method": "openpyxl_read_only_source", "page_type": "worksheet_row",
                    "text_quality": "unverified", "char_count": len(text), "image_count": 0, "error": None}
            rows.append(item)
            state["evidence"][eid] = item
            state["sections"][secid]["evidence_ids"].append(eid)
        sheets.append({"sheet_name": ws.title, "state": ws.sheet_state, "dimension": ws.calculate_dimension(),
                       "formulas": formulas, "merged_ranges": [str(x) for x in ws.merged_cells.ranges],
                       "hidden_rows": [r for r, x in ws.row_dimensions.items() if x.hidden],
                       "hidden_columns": [c for c, x in ws.column_dimensions.items() if x.hidden],
                       "protected": bool(ws.protection.sheet), "tables": list(ws.tables),
                       "image_count": len(ws._images), "chart_count": len(ws._charts)})
        if ws._images or ws._charts:
            add_issue(state, "workbook_visual_review", secid, f"{ws.title}: embedded images/charts need visual review.")
    doc.update(sheet_count=len(sheets), sheets=sheets, defined_names=[str(n) for n in wb.defined_names.values()],
               external_links=len(wb._external_links), extraction_status="extracted_unverified")
    if wb._external_links:
        add_issue(state, "external_workbook_link", doc["document_id"], "External workbook dependencies require review.")
    jsonl(root / "cells.jsonl", cells)
    jsonl(root / "rows.jsonl", rows)
    dump(root / "workbook.json", sheets)
    wb.close()
    cached.close()


def ingest(args):
    project, source = Path(args.project).resolve(), Path(args.source).resolve()
    if not source.is_dir() or source == project or source.is_relative_to(project):
        raise ValueError("Source must be an existing folder outside the project packet")
    paths = scan_source(source, project, args.exclude)
    with project_lock(project):
        statepath = project / INTAKE / "state.json"
        state = load(project) if statepath.exists() else blank_state(args.name or source.name)
        old_snapshot, old_files = state["snapshot"], state["files"]
        previous_docs = active_docs(state)
        files, processed, reused = [], 0, 0
        for path in paths:
            rel = rel_path(path, source)
            if path.is_symlink():
                files.append({"relative_path": rel, "sha256": None, "document_id": None, "status": "symlink_skipped"})
                continue
            h = sha(path)
            did = "DOC-" + h[:16]
            entry = {"relative_path": rel, "sha256": h, "document_id": did, "bytes": path.stat().st_size,
                     "original_absolute_path": str(path), "status": "present"}
            files.append(entry)
            if did in state["documents"]:
                if state["documents"][did]["sha256"] != h:
                    raise ValueError("Document identifier collision")
                reused += 1
                continue
            dest = project / SOURCE / did / path.name
            dest.parent.mkdir(parents=True, exist_ok=True)
            if dest.exists() and sha(dest) != h:
                raise ValueError(f"Source-copy integrity failure at {dest}")
            if not dest.exists():
                shutil.copy2(path, dest)
            if sha(dest) != h or sha(path) != h:
                raise ValueError(f"Source changed while copying {path}")
            doc = {"document_id": did, "sha256": h, "original_filename": path.name,
                   "source_path": rel_path(dest, project), "bytes": entry["bytes"],
                   "file_type": path.suffix.lower(), "role": classify_filename(path.name),
                   "role_status": "filename_candidate", "first_received_at": now(),
                   "authority_status": "unverified", "extraction_version": VERSION}
            state["documents"][did] = doc
            print(f"Processing {path.name}", file=sys.stderr, flush=True)
            try:
                if path.suffix.lower() == ".pdf":
                    extract_pdf(project, state, doc)
                elif path.suffix.lower() == ".xlsx":
                    extract_workbook(project, state, doc)
                else:
                    doc["extraction_status"] = "unsupported"
                    add_issue(state, "unsupported_file", did, f"No extractor for {path.suffix or 'extensionless file'}")
            except Exception as exc:
                doc["extraction_status"] = "failed"
                add_issue(state, "document_extraction_error", did, str(exc))
            # Integrity manifest for every extraction artifact, including cell/word/table data.
            extraction_root = project / EXTRACT / did
            doc["extraction_hashes"] = {rel_path(p, project): sha(p) for p in extraction_root.rglob("*") if p.is_file()} if extraction_root.exists() else {}
            processed += 1
        state["files"] = files
        refine_sections(state)
        state["source_folder"] = str(source)
        state["source_excludes"] = args.exclude
        state["snapshot"] = "SNAP-" + digest([(x["relative_path"], x["sha256"]) for x in files])[:20]
        current_docs = active_docs(state)
        if old_snapshot != state["snapshot"]:
            changed_docs = previous_docs ^ current_docs
            affected = {k for k, r in state["records"].items() if any(state["evidence"][e]["document_id"] in changed_docs for e in r["evidence_ids"])}
            # Candidate amendment impacts are conservative: every prior record needs review until mapped.
            new_amendment = any(state["documents"][d]["role"] == "amendment_candidate" for d in current_docs - previous_docs)
            if new_amendment:
                affected.update(state["records"])
            grew = True
            while grew:
                before = len(affected)
                for edge in state["relationships"].values():
                    if edge["from_id"] in affected or edge["to_id"] in affected:
                        affected.update([edge["from_id"], edge["to_id"]])
                grew = len(affected) > before
            old_by_path = {f["relative_path"]: f for f in old_files}
            replacements = []
            for f in files:
                old = old_by_path.get(f["relative_path"])
                if old and old["sha256"] != f["sha256"]:
                    replacements.append({"path": f["relative_path"], "old_document_id": old["document_id"],
                                         "new_document_id": f["document_id"], "effect": "candidate_replacement_not_supersession"})
                    a = "\n".join(e["text"] for e in state["evidence"].values() if e["document_id"] == old["document_id"])
                    b = "\n".join(e["text"] for e in state["evidence"].values() if e["document_id"] == f["document_id"])
                    write_text(project / INTAKE / "changes" / f"{state['snapshot']}-{f['document_id']}.diff",
                               "\n".join(difflib.unified_diff(a.splitlines(), b.splitlines(), fromfile=old["document_id"], tofile=f["document_id"])))
            report = {"snapshot": state["snapshot"], "previous_snapshot": old_snapshot, "created_at": now(),
                      "added_documents": sorted(current_docs - previous_docs), "removed_documents": sorted(previous_docs - current_docs),
                      "candidate_replacements": replacements, "affected_records": sorted(affected),
                      "estimate_lines_requiring_review": [x for x in state["estimate_links"].values() if x["record_id"] in affected],
                      "all_approvals_invalidated": bool(old_snapshot), "new_amendment_candidate": new_amendment}
            state["change_reports"].append(report)
            dump(project / INTAKE / "changes" / (state["snapshot"] + ".json"), report)
            dump(project / INTAKE / "snapshots" / (state["snapshot"] + ".json"), {"snapshot": state["snapshot"], "files": files, "at": now()})
        save(project, state)
        publish(project, state)
        print(json.dumps({"snapshot": state["snapshot"], "physical_files": len(files), "unique_documents": len(current_docs),
                          "new_documents_extracted": processed, "unchanged_or_duplicate_files": reused, "start_here": str(project / "START_HERE.md")}))


def validate_schema(value, schema, root, path="batch"):
    """Validate the small standard JSON Schema subset used in batch.schema.json."""
    if "$ref" in schema:
        target = root
        for part in schema["$ref"].split("/")[1:]:
            target = target[part]
        return validate_schema(value, target, root, path)
    typ = schema.get("type")
    tests = {"object": lambda v: isinstance(v, dict), "array": lambda v: isinstance(v, list),
             "string": lambda v: isinstance(v, str), "boolean": lambda v: type(v) is bool,
             "integer": lambda v: type(v) is int, "number": lambda v: type(v) in (int, float), "null": lambda v: v is None}
    if typ and not any(tests[t](value) for t in (typ if isinstance(typ, list) else [typ])):
        raise ValueError(f"{path}: expected {typ}")
    if "enum" in schema and value not in schema["enum"]:
        raise ValueError(f"{path}: invalid value {value!r}")
    if isinstance(value, str):
        if len(value.strip()) < schema.get("minLength", 0):
            raise ValueError(f"{path}: text is required")
        if "pattern" in schema and not re.search(schema["pattern"], value):
            raise ValueError(f"{path}: invalid identifier")
    if isinstance(value, list):
        if len(value) < schema.get("minItems", 0):
            raise ValueError(f"{path}: at least {schema['minItems']} item(s) required")
        for i, v in enumerate(value):
            validate_schema(v, schema.get("items", {}), root, f"{path}[{i}]")
    if isinstance(value, dict):
        for key in schema.get("required", []):
            if key not in value:
                raise ValueError(f"{path}: missing {key}")
        props = schema.get("properties", {})
        if schema.get("additionalProperties") is False:
            extra = set(value) - set(props)
            if extra:
                raise ValueError(f"{path}: unsupported fields {sorted(extra)}")
        for k, v in value.items():
            if k in props:
                validate_schema(v, props[k], root, path + "." + k)


def require_evidence(state, ids):
    if not ids or any(e not in state["evidence"] for e in ids):
        raise ValueError("Every evidence reference must resolve; at least one is required")


def review_token(state):
    # Exclude the attestations themselves; any semantic/decision/source change invalidates them.
    return digest({k: state[k] for k in ("snapshot", "records", "relationships", "expected_documents", "issues", "section_reviews", "ocr", "evidence")})


def import_batch(args):
    project = Path(args.project).resolve()
    batch = read_json(args.batch)
    schema = read_json(Path(__file__).parent.parent / "references" / "batch.schema.json")
    validate_schema(batch, schema, schema)
    with project_lock(project):
        state = load(project)
        if batch["snapshot"] != state["snapshot"]:
            raise ValueError("Batch snapshot is stale; inspect current changes before importing")
        bid, bh = batch["batch_id"], digest(batch)
        if bid in state["batches"]:
            if state["batches"][bid] != bh:
                raise ValueError("Batch ID is immutable; use a new ID for a revision")
            print("Batch already imported; no changes")
            return
        for rec in batch.get("records", []):
            if rec["authority_status"] == "human_verified":
                raise ValueError("Human authority is established through reviewer decisions, not a record label")
            require_evidence(state, rec["evidence_ids"])
            if not rec["record_id"].startswith(TYPES[rec["type"]] + "-"):
                raise ValueError(f"Record ID prefix must match type: {rec['record_id']}")
            for citation in rec["citations"]:
                eid = citation["evidence_id"]
                if eid not in rec["evidence_ids"]:
                    raise ValueError("Citation must be included in evidence_ids")
                if norm(citation["quote"]) not in norm(state["evidence"][eid]["text"]):
                    raise ValueError(f"Quote does not match extracted evidence {eid}: {citation['quote'][:80]}")
            if set(rec["evidence_ids"]) != {c["evidence_id"] for c in rec["citations"]}:
                raise ValueError("Each evidence reference requires a supporting quote")
            old = state["records"].get(rec["record_id"])
            if old and digest(old) != digest(rec):
                raise ValueError("Records are immutable; create a new ID and a version relationship")
            state["records"][rec["record_id"]] = rec
        for edge in batch.get("relationships", []):
            if edge["from_id"] not in state["records"] or edge["to_id"] not in state["records"]:
                raise ValueError("Relationship endpoints must be records")
            if edge["from_id"] == edge["to_id"]:
                raise ValueError("Self relationships are not allowed")
            require_evidence(state, edge["evidence_ids"])
            if edge["relation"] == "supersedes" and (not edge.get("human_approved") or not edge.get("reviewer") or not edge.get("decision_note")):
                raise ValueError("Supersession requires explicit human approval, reviewer, note, and source evidence")
            key = edge["relationship_id"]
            if key in state["relationships"] and state["relationships"][key] != edge:
                raise ValueError("Relationships are immutable; use a new identifier")
            state["relationships"][key] = edge
        # An approved supersession graph must remain acyclic.
        supersession = collections.defaultdict(list)
        for edge in state["relationships"].values():
            if edge["relation"] == "supersedes":
                supersession[edge["from_id"]].append(edge["to_id"])
        def check_cycle(node, visiting, visited):
            if node in visiting:
                raise ValueError("Supersession cycle detected")
            if node in visited:
                return
            visiting.add(node)
            for child in supersession.get(node, []):
                check_cycle(child, visiting, visited)
            visiting.remove(node)
            visited.add(node)
        visited = set()
        for node in list(supersession):
            check_cycle(node, set(), visited)
        for expected in batch.get("expected_documents", []):
            require_evidence(state, expected["evidence_ids"])
            for did in expected.get("matched_document_ids", []):
                if did not in state["documents"]:
                    raise ValueError("Matched expected document does not exist")
            state["expected_documents"][expected["expected_id"]] = expected
        for issue in batch.get("issues", []):
            require_evidence(state, issue["evidence_ids"])
            key = issue["issue_id"]
            if key in state["issues"] and state["issues"][key] != issue:
                raise ValueError("Issue ID exists; add a resolution rather than replacing it")
            state["issues"][key] = issue
        for res in batch.get("resolutions", []):
            issue = state["issues"].get(res["issue_id"])
            if not issue:
                raise ValueError("Resolution points to unknown issue")
            require_evidence(state, res["evidence_ids"])
            issue["resolution"] = dict(res, snapshot=state["snapshot"], at=now())
        for review in batch.get("section_reviews", []):
            section = state["sections"].get(review["section_id"])
            if not section or set(review["covered_evidence_ids"]) != set(section["evidence_ids"]):
                raise ValueError("Section review must cover every evidence unit in that section")
            state["section_reviews"][review["section_id"]] = dict(review, snapshot=state["snapshot"])
        for link in batch.get("estimate_links", []):
            if link["record_id"] not in state["records"]:
                raise ValueError("Estimate link record does not exist")
            state["estimate_links"][link["estimate_line_id"] + ":" + link["record_id"]] = link
        for approval in batch.get("attestations", []):
            state["attestations"][approval["kind"]] = dict(approval, snapshot=state["snapshot"],
                                                          review_token=review_token(state), at=now())
        state["batches"][bid] = bh
        dump(project / INTAKE / "review_batches" / (bid + ".json"), batch)
        save(project, state)
        publish(project, state)
        print(json.dumps({"imported": bid, "records_in_batch": len(batch.get('records', [])),
                          "state_records": len(state["records"]), "status": compute_gate(project, state)["status"]}))


def compute_gate(project, state):
    docs = active_docs(state)
    blockers = []
    def block(kind, target, message):
        blockers.append({"kind": kind, "target": target, "message": message})
    if not state["files"]:
        block("empty_package", "package", "No input files")
    # A live input folder changing must block even before the operator reruns ingest.
    upstream = Path(state.get("source_folder", ""))
    if state.get("source_folder") and upstream.is_dir():
        observed_input = [(rel_path(p, upstream), None if p.is_symlink() else sha(p))
                          for p in scan_source(upstream, project, state.get("source_excludes", []))]
        saved_input = [(f["relative_path"], f["sha256"]) for f in state["files"]]
        if observed_input != saved_input:
            block("source_folder_changed", "package", "Input folder changed since intake. Run ingest against the complete current package.")
    for f in state["files"]:
        if not f.get("document_id"):
            block("unsupported_source", f["relative_path"], f["status"])
    for did in docs:
        doc = state["documents"][did]
        path = project / doc["source_path"]
        if not path.exists() or sha(path) != doc["sha256"]:
            block("source_integrity", did, "Source-of-record copy is missing or altered")
        if doc["extraction_status"] != "extracted_unverified":
            block("extraction_failure", did, doc["extraction_status"])
        for rel, expected_hash in doc.get("extraction_hashes", {}).items():
            p = project / rel
            if not p.exists() or sha(p) != expected_hash:
                block("extraction_integrity", did, f"Extracted artifact missing or altered: {rel}")
        if doc.get("page_count") is not None:
            observed = [e["pdf_page"] for e in state["evidence"].values() if e["document_id"] == did and e["method"] == "pdfplumber_native"]
            if sorted(observed) != list(range(1, doc["page_count"] + 1)):
                block("page_coverage", did, "Not every PDF page has exactly one evidence record")
    for expected in state["expected_documents"].values():
        matched = set(expected.get("matched_document_ids", [])) & docs
        if not matched:
            block("missing_document", expected["expected_id"], expected["title"])
        elif not expected.get("version_verified"):
            block("version_unverified", expected["expected_id"], expected["title"])
    if not state["expected_documents"]:
        block("inventory_unreviewed", "package", "Expected attachments/references have not been reconciled")
    for sid, sec in state["sections"].items():
        if sec["document_id"] not in docs:
            continue
        review = state["section_reviews"].get(sid, {})
        if review.get("snapshot") != state["snapshot"] or set(review.get("covered_evidence_ids", [])) != set(sec["evidence_ids"]):
            block("semantic_review", sid, sec["heading"])
    for item in state["ocr"].values():
        path = project / item["path"] / "ocr.txt"
        if not path.exists() or sha(path) != item["text_sha256"]:
            block("ocr_integrity", item["path"], "OCR sidecar is missing or changed")
    for e in state["evidence"].values():
        if e["method"] == "tesseract_ocr_unverified":
            p = project / e["text_path"]
            if not p.exists() or sha(p) != e["text_file_sha256"]:
                block("ocr_integrity", e["evidence_id"], "Referenced OCR artifact is missing or changed")
    for key, issue in state["issues"].items():
        if issue.get("severity", "blocking") != "blocking":
            continue
        target = issue.get("target_id")
        if target in state["documents"] and target not in docs:
            continue
        if target in state["evidence"] and state["evidence"][target]["document_id"] not in docs:
            continue
        res = issue.get("resolution", {})
        if res.get("snapshot") != state["snapshot"] or not res.get("human_approved"):
            block(issue["kind"], key, issue["message"])
    # Candidate supersession and conflicts block until a documented issue resolution covers the edge.
    for edge in state["relationships"].values():
        if edge["relation"] in {"candidate_supersedes", "conflicts_with"}:
            resolved = any(i.get("target_id") == edge["relationship_id"] and
                           i.get("resolution", {}).get("human_approved") and
                           i.get("resolution", {}).get("snapshot") == state["snapshot"] for i in state["issues"].values())
            if not resolved:
                block("authority_conflict", edge["relationship_id"], f"{edge['from_id']} {edge['relation']} {edge['to_id']}")
    token = review_token(state)
    for kind in sorted(ATTESTATIONS):
        att = state["attestations"].get(kind, {})
        if att.get("snapshot") != state["snapshot"] or att.get("review_token") != token or not att.get("human_approved"):
            block("human_review", kind, "Named human approval absent or stale")
    counts = collections.Counter(x["kind"] for x in blockers)
    status = "READY" if not blockers else ("BLOCKED_MISSING_SOURCE" if counts["missing_document"] else "BLOCKED_REVIEW")
    return {"schema_version": VERSION, "snapshot": state["snapshot"], "review_token": token, "status": status,
            "estimating_allowed": not blockers, "checked_at": now(), "physical_files": len(state["files"]),
            "unique_documents": len(docs), "pdf_pages": sum(state["documents"][d].get("page_count", 0) for d in docs),
            "expected_document_count": len(state["expected_documents"]), "blocker_count": len(blockers),
            "blockers_by_kind": dict(counts), "blockers": blockers}


def md_escape(text):
    return str(text).replace("|", "\\|").replace("\n", " ")


def evidence_link(project, state, evid, from_path=None):
    e = state["evidence"][evid]
    doc = state["documents"][e["document_id"]]
    target = project / doc["source_path"]
    label = doc["original_filename"]
    suffix = ""
    if e.get("pdf_page"):
        label += f", PDF p. {e['pdf_page']}"
        suffix = f"#page={e['pdf_page']}"
    else:
        label += f", {e['sheet_name']}!{e['cell_or_range']}"
    relative = os.path.relpath(target, (from_path or (project / "START_HERE.md")).parent)
    return f"[{md_escape(label)}]({quote(relative, safe='/')}" + suffix + ")"


def publish(project, state):
    docs = active_docs(state)
    gate = compute_gate(project, state)
    dump(project / INTAKE / "intake_gate.json", gate)
    csv_write(project / INTAKE / "document_manifest.csv", state["files"],
              ["relative_path", "document_id", "sha256", "bytes", "status", "original_absolute_path"])
    csv_write(project / INTAKE / "extraction_status.csv", state["documents"].values(),
              ["document_id", "original_filename", "file_type", "role", "page_count", "sheet_count", "extraction_status", "classification_counts"])
    csv_write(project / INTAKE / "document_versions.csv", [dict(d, in_current_snapshot=did in docs) for did,d in state["documents"].items()],
              ["document_id", "sha256", "original_filename", "in_current_snapshot", "first_received_at", "authority_status", "source_path", "extraction_version"])
    duplicates = []
    for did in docs:
        aliases = [f["relative_path"] for f in state["files"] if f.get("document_id") == did]
        if len(aliases) > 1:
            duplicates.append({"document_id": did, "aliases": aliases, "physical_count": len(aliases), "extraction_count": 1})
    csv_write(project / INTAKE / "duplicate_files.csv", duplicates, ["document_id", "aliases", "physical_count", "extraction_count"])
    table_register = []
    for did in docs:
        path = project / EXTRACT / did / "tables.jsonl"
        if path.exists():
            for line in path.read_text(encoding="utf-8").splitlines():
                t = json.loads(line)
                table_register.append({"table_id":t["table_id"],"evidence_id":t["evidence_id"],"document_id":did,
                                       "pdf_page":t["pdf_page"],"bbox":t["bbox"],"row_count":len(t["rows"]),
                                       "empty_cells":sum(v is None for row in t["rows"] for v in row),
                                       "review_status":t["review_status"],"table_data_path":rel_path(path, project)})
    csv_write(project / EXTRACT / "table_extraction_register.csv", table_register,
              ["table_id","document_id","evidence_id","pdf_page","bbox","row_count","empty_cells","review_status","table_data_path"])
    jsonl(project / INTAKE / "processing_log.jsonl", [{"event":"snapshot","data":r} for r in state["change_reports"]] +
          [{"event":"review_batch","batch_id":bid,"sha256":h,"path":f"review_batches/{bid}.json"} for bid,h in state["batches"].items()])
    csv_write(project / INTAKE / "expected_documents.csv", state["expected_documents"].values(),
              ["expected_id", "title", "matched_document_ids", "version_verified", "evidence_ids", "note"])
    csv_write(project / INTAKE / "review_queue.csv", gate["blockers"], ["kind", "target", "message"])
    csv_write(project / EXTRACT / "page_reference_map.csv", active_evidence(state).values(),
              ["evidence_id", "document_id", "source_sha256", "pdf_page", "printed_page", "printed_page_scheme",
               "section_number", "heading", "section_id", "sheet_name", "cell_or_range", "page_type", "method", "text_quality"])
    # Build complete structural section files; indexing preserves evidence granularity.
    section_index = []
    for sid, sec in state["sections"].items():
        path = project / EXTRACT / "sections" / (sid + ".md")
        sec["path"] = rel_path(path, project)
        lines = ["# " + sec["heading"], "", "Faithful extracted text. The linked original is authoritative.", "",
                 "Section boundary: " + sec["method"] + ". Review before treating it as a complete semantic unit.", ""]
        for eid in sec["evidence_ids"]:
            e = state["evidence"][eid]
            lines += [f"## {eid}", "", evidence_link(project, state, eid, path), "", "```text", e["text"], "```", ""]
        write_text(path, "\n".join(lines))
        if sec["document_id"] in docs:
            section_index.append(sec)
    csv_write(project / INTAKE / "semantic_sections.csv", section_index, ["section_id", "document_id", "heading", "path", "evidence_ids", "method"])
    # Canonical JSON remains typed; CSVs and wiki pages are generated views.
    jsonl(project / KNOWLEDGE / "records.jsonl", state["records"].values())
    jsonl(project / KNOWLEDGE / "relationships.jsonl", state["relationships"].values())
    csv_write(project / KNOWLEDGE / "relationships.csv", state["relationships"].values(),
              ["relationship_id", "from_id", "relation", "to_id", "evidence_ids", "human_approved", "reviewer", "decision_note"])
    for typ, filename in FILES.items():
        records = [r for r in state["records"].values() if r["type"] == typ]
        csv_write(project / KNOWLEDGE / (filename + ".csv"), records,
                  ["record_id", "title", "interpretation", "contract_layer", "applicability", "tags", "critical",
                   "authority_status", "review_status", "evidence_ids", "citations", "details"])
        index = ["# " + filename.replace("_", " ").title(), "", f"{len(records)} records. AI interpretations require review.", ""]
        for r in records:
            index.append(f"- [{r['record_id']}: {md_escape(r['title'])}](records/{r['record_id']}.md)")
            path = project / KNOWLEDGE / "records" / (r["record_id"] + ".md")
            lines = [f"# {r['record_id']}: {r['title']}", "", f"Layer: {r['contract_layer']}. Authority: {r['authority_status']}. Review: {r['review_status']}.",
                     "", "## AI interpretation", "", r["interpretation"], "", "## Source evidence", ""]
            for cite in r["citations"]:
                lines += [evidence_link(project, state, cite["evidence_id"], path), "", "> " + cite["quote"].replace("\n", "\n> "), ""]
            if r.get("details"):
                lines += ["## Structured fields", "", "```json", json.dumps(r["details"], indent=2, ensure_ascii=False), "```", ""]
            lines += ["## Related records", ""]
            for edge in state["relationships"].values():
                if r["record_id"] in (edge["from_id"], edge["to_id"]):
                    other = edge["to_id"] if edge["from_id"] == r["record_id"] else edge["from_id"]
                    lines.append(f"- {edge['relation']}: [{other}]({other}.md) ({edge['relationship_id']})")
            write_text(path, "\n".join(lines))
        write_text(project / KNOWLEDGE / (filename + ".md"), "\n".join(index) + "\n")
    # Local FTS is a disposable index, never the source of authority.
    dbpath = project / INTAKE / "search.sqlite"
    with contextlib.closing(sqlite3.connect(dbpath)) as con, con:
        con.execute("DROP TABLE IF EXISTS evidence_search")
        con.execute("CREATE VIRTUAL TABLE evidence_search USING fts5(evidence_id UNINDEXED, document_id UNINDEXED, heading, text)")
        con.executemany("INSERT INTO evidence_search VALUES(?,?,?,?)",
                        [(e["evidence_id"], e["document_id"], e["heading"], e["text"]) for e in active_evidence(state).values()])
    # Save section paths before callers reload state.
    save(project, state)
    lines = ["# " + state["project_name"], "", "## Intake status", "",
             f"**{gate['status']}** — estimating allowed: **{str(gate['estimating_allowed']).lower()}**.", "",
             f"{len(state['files'])} physical files, {len(docs)} unique documents, {gate['pdf_pages']} PDF pages, "
             f"{len(state['records'])} registered records, {len(section_index)} sections to review.", "",
             "The registers contain draft interpretations. Extraction coverage is separate from semantic completeness.", "",
             "[Readiness details](01_INTAKE/intake_gate.json) · [Review queue](01_INTAKE/review_queue.csv) · "
             "[Expected documents](01_INTAKE/expected_documents.csv)", "", "## Project registers", ""]
    for filename in FILES.values():
        lines.append(f"- [{filename.replace('_', ' ').title()}](03_PROJECT%20KNOWLEDGE/{filename}.md)")
    lines += ["", "## Source documents and sections", ""]
    for did in sorted(docs):
        doc = state["documents"][did]
        aliases = [f["relative_path"] for f in state["files"] if f.get("document_id") == did]
        lines += [f"### {md_escape(doc['original_filename'])}", "",
                  f"{did} · {doc['role']} · {doc['extraction_status']}", "",
                  f"[Original file]({quote(doc['source_path'], safe='/')})", ""]
        if len(aliases) > 1:
            lines += ["Exact duplicate aliases: " + "; ".join(aliases), ""]
        for sec in section_index:
            if sec["document_id"] == did:
                lines.append(f"- [{md_escape(sec['heading'])}]({quote(sec['path'], safe='/')}) ({len(sec['evidence_ids'])} evidence units)")
        lines += [""]
    lines += ["## Current blocker counts", "", "| Check | Count |", "|---|---:|"]
    for key, count in sorted(gate["blockers_by_kind"].items()):
        lines.append(f"| {key} | {count} |")
    write_text(project / "START_HERE.md", "\n".join(lines) + "\n")
    write_text(project / KNOWLEDGE / "project_brief.md", "\n".join([
        "# " + state["project_name"], "", "AI-generated orientation; source files and review status control use.", "",
        f"Current status: {gate['status']}. Estimating allowed: {str(gate['estimating_allowed']).lower()}.", "",
        "[Open project index](../START_HERE.md) · [Scope](scope_register.md) · [Forms](forms_checklist.md) · [Amendments](amendment_log.md)", "",
        "## Common constraints", ""] +
        [f"- [{r['record_id']}: {md_escape(r['title'])}](records/{r['record_id']}.md) — {r['interpretation']}"
         for r in state["records"].values() if r.get("always_include")] + ["",
        "The source set is indexed. The knowledge registers remain draft until section review and named human decisions are complete.", ""]))
    dump(project / KNOWLEDGE / "estimator_context_index.json", {
        "snapshot": state["snapshot"], "start_here": "START_HERE.md", "gate": INTAKE + "/intake_gate.json",
        "common_record_ids": [r["record_id"] for r in state["records"].values() if r.get("always_include")],
        "records": KNOWLEDGE + "/records.jsonl", "relationships": KNOWLEDGE + "/relationships.jsonl",
        "retrieval": "Use rfp_intake.py packet; check live gate before estimating. Keywords alone are not a completeness check."})
    for d in ("04_ESTIMATING", "05_FINAL RESPONSE"):
        (project / d).mkdir(exist_ok=True)


def search(project, state, query, limit):
    terms = re.findall(r"[\w]+", query)
    if not terms:
        return []
    fts = " OR ".join('"' + t.replace('"', '""') + '"' for t in terms)
    with contextlib.closing(sqlite3.connect(project / INTAKE / "search.sqlite")) as con, con:
        matches = con.execute("SELECT evidence_id FROM evidence_search WHERE evidence_search MATCH ? ORDER BY bm25(evidence_search) LIMIT ?", (fts, limit)).fetchall()
    return [state["evidence"][x[0]] for x in matches]


def wire_json(value):
    """The exact compact serialization used for model-facing packets."""
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def measure_packet(result):
    # Include the measurements themselves in the measured payload.
    result["usage"] = {"serialized_characters": 0, "utf8_bytes": 0,
                       "estimated_tokens": 0, "method": "ceil(UTF-8 bytes / 4); heuristic, not tokenizer or billing"}
    for _ in range(12):
        payload = wire_json(result) + "\n"
        sizes = (len(payload), len(payload.encode("utf-8")))
        usage = result["usage"]
        if (usage["serialized_characters"], usage["utf8_bytes"]) == sizes:
            return payload
        usage.update(serialized_characters=sizes[0], utf8_bytes=sizes[1], estimated_tokens=(sizes[1] + 3) // 4)
    raise ValueError("Packet measurement did not stabilize")


def page_context(base, collections, fingerprint, cursor, max_chars):
    """Lossless paging: oversized individual objects become explicitly labeled JSON fragments."""
    if not 4000 <= max_chars <= 100000:
        raise ValueError("--max-chars must be between 4000 and 100000")
    units = []
    # Reserve room for metadata and JSON escaping of a fragment inside JSON.
    fragment_size = min(4000, max_chars // 8)
    for collection, values in collections.items():
        for value in values:
            encoded = wire_json(value)
            if len(encoded) <= fragment_size:
                units.append((collection, value))
            else:
                item_id = identifier("ITEM", [collection, value])
                chunks = [encoded[i:i+fragment_size] for i in range(0, len(encoded), fragment_size)]
                for index, chunk in enumerate(chunks):
                    units.append(("fragments", {"collection": collection, "item_id": item_id,
                        "part": index + 1, "parts": len(chunks), "json_fragment": chunk}))
    offset = 0
    if cursor:
        try:
            token, raw_offset = cursor.rsplit(":", 1)
            offset = int(raw_offset)
        except (ValueError, AttributeError):
            raise ValueError("Invalid cursor")
        if token != fingerprint:
            raise ValueError("Stale cursor: package, review, gate, query, view, or budget changed; restart retrieval")
        if offset < 0 or offset >= len(units):
            raise ValueError("Cursor offset outside this result")
    result = dict(base, **{key: [] for key in collections}, fragments=[])
    result["pagination"] = {"has_more": False, "next_cursor": None, "start_unit": offset,
                            "returned_units": 0, "total_units": len(units), "max_characters": max_chars}
    result["fragment_instructions"] = "Join json_fragment parts in part order for each item_id, then parse JSON; never interpret an incomplete item."

    def update_page(end):
        result["pagination"].update(has_more=end < len(units),
            next_cursor=f"{fingerprint}:{end}" if end < len(units) else None, returned_units=end-offset)
        return measure_packet(result)

    payload = update_page(offset)
    if len(payload) > max_chars:
        raise ValueError("Packet metadata exceeds budget; shorten the query or increase --max-chars")
    end = offset
    for collection, value in units[offset:]:
        result[collection].append(value)
        candidate = update_page(end + 1)
        if len(candidate) > max_chars:
            result[collection].pop()
            payload = update_page(end)
            if end == offset:
                raise ValueError("Single fragment exceeds budget; increase --max-chars")
            break
        payload = candidate
        end += 1
    return result, payload


def packet(args):
    project = Path(args.project).resolve()
    state = load(project)
    gate = compute_gate(project, state)
    if args.mode == "estimating" and not gate["estimating_allowed"]:
        print(wire_json({"error": "Estimating is blocked", "gate_status": gate["status"],
            "blockers_by_kind": gate["blockers_by_kind"],
            "next_step": "Use packet --mode research --view blockers for paginated blocker details"}))
        return 2
    view = getattr(args, "view", "summary")
    max_chars = getattr(args, "max_chars", 24000)
    cursor = getattr(args, "cursor", None)
    record_ids = set(getattr(args, "record_id", None) or [])
    evidence_ids = set(getattr(args, "evidence_id", None) or [])
    if args.limit < 1:
        raise ValueError("--limit must be positive")
    if record_ids - state["records"].keys() or evidence_ids - state["evidence"].keys():
        raise ValueError("Unknown record or evidence ID")
    if not args.query.strip() and not record_ids and not evidence_ids and view != "blockers":
        raise ValueError("Provide --query, --record-id, or --evidence-id")
    matches = search(project, state, args.query, args.limit) if not record_ids and not evidence_ids else []
    eids = {e["evidence_id"] for e in matches}
    eids.update(evidence_ids)
    terms = set(re.findall(r"\w+", args.query.lower())) if not record_ids and not evidence_ids else set()
    selected = set(record_ids)
    for rid, r in state["records"].items():
        text = " ".join([r["title"], r["interpretation"], " ".join(r["tags"])]).lower()
        if r.get("always_include") or eids.intersection(r["evidence_ids"]) or terms.intersection(re.findall(r"\w+", text)):
            selected.add(rid)
    # Include full connected components. Report size; do not silently discard conflicting records.
    while True:
        before = len(selected)
        for edge in state["relationships"].values():
            if edge["from_id"] in selected or edge["to_id"] in selected:
                selected.update([edge["from_id"], edge["to_id"]])
        if len(selected) == before:
            break
    for rid in selected:
        eids.update(state["records"][rid]["evidence_ids"])
    evidence = []
    for eid in sorted(eids):
        e = dict(state["evidence"][eid])
        e["source_path"] = state["documents"][e["document_id"]]["source_path"]
        e["in_current_snapshot"] = e["document_id"] in active_docs(state)
        evidence.append(e)
    related_edges = [e for e in state["relationships"].values() if e["from_id"] in selected or e["to_id"] in selected]
    records = [state["records"][r] for r in sorted(selected)]
    all_result = {"snapshot": state["snapshot"], "mode": args.mode, "query": args.query,
              "estimating_allowed": gate["estimating_allowed"], "gate_status": gate["status"],
              "blockers": gate["blockers"], "retrieval_method": "local_full_text_plus_relationships",
              "completeness_claim": False, "records": records,
              "relationships": related_edges,
              "evidence": evidence, "approximate_characters": sum(len(e["text"]) for e in evidence)}
    fingerprint = digest([all_result, gate["review_token"], view, max_chars, args.limit, sorted(record_ids), sorted(evidence_ids)])
    # Full context stays on disk, not in the model response. Never overwrite a prior context artifact.
    archive = project / "04_ESTIMATING" / "context-cache" / (fingerprint + ".json")
    if archive.exists():
        if read_json(archive) != all_result:
            raise ValueError("Cached context was modified; refuse to use it")
    else:
        dump(archive, all_result)
    conflicts = [e for e in related_edges if e["relation"] in {"conflicts_with", "candidate_supersedes"}]
    common_ids = {r["record_id"] for r in records if r.get("always_include")}
    priority_ids = common_ids | {rid for e in conflicts for rid in (e["from_id"], e["to_id"])}
    records.sort(key=lambda r: (r["record_id"] not in priority_ids, r["record_id"]))
    targets = selected | {e["relationship_id"] for e in related_edges} | eids
    related_blockers = [b for b in gate["blockers"] if b["target"] in targets]
    base = {k: all_result[k] for k in ("snapshot", "mode", "query", "estimating_allowed", "gate_status", "completeness_claim")}
    base.update(format_version=2, view=view, context_id=fingerprint,
        full_context_file=rel_path(archive, project),
        blockers_by_kind=gate["blockers_by_kind"],
        totals={"records": len(records), "relationships": len(related_edges), "evidence": len(evidence),
                "blockers": len(gate["blockers"]), "related_blockers": len(related_blockers),
                "common_records": len(common_ids), "conflict_relationships": len(conflicts)},
        warnings=["Retrieval is not a completeness check. Summary cards omit evidence text and record details.",
                  "Common constraints and connected conflicts may continue on later pages; follow next_cursor before relying on this context.",
                  "Use --view full with --record-id or --evidence-id for exact details; --view blockers lists all blockers."])
    if view == "summary":
        cards = [{k: r[k] for k in ("record_id", "type", "title", "interpretation", "applicability", "contract_layer",
                    "authority_status", "review_status", "critical", "always_include", "evidence_ids")} for r in records]
        locators = [{k: e[k] for k in ("evidence_id", "document_id", "source_path", "source_sha256", "pdf_page", "printed_page",
                    "sheet_name", "cell_or_range", "text_path", "section_id", "parent_evidence_id", "bbox", "method", "in_current_snapshot") if k in e} for e in evidence]
        collections = {"relationships": related_edges, "records": cards, "evidence": locators, "blockers": related_blockers}
    elif view == "full":
        collections = {"relationships": related_edges, "records": records, "evidence": evidence, "blockers": related_blockers}
    else:
        collections = {"blockers": gate["blockers"], "relationships": [], "records": [], "evidence": []}
    result, payload = page_context(base, collections, fingerprint, cursor, max_chars)
    if args.output:
        out = Path(args.output).resolve()
        if not out.is_relative_to(project / "04_ESTIMATING"):
            raise ValueError("Packet output must be inside project/04_ESTIMATING")
        write_text(out, payload)
        print(wire_json({"output": str(out), "totals": result["totals"], "pagination": result["pagination"],
                        "usage": result["usage"], "estimating_allowed": result["estimating_allowed"]}))
    else:
        print(payload, end="")
    return 0


def next_section(args):
    project = Path(args.project).resolve()
    state = load(project)
    for sid, sec in state["sections"].items():
        if sec["document_id"] in active_docs(state) and state["section_reviews"].get(sid, {}).get("snapshot") != state["snapshot"]:
            print(json.dumps(dict(sec, absolute_path=str(project / sec["path"]), snapshot=state["snapshot"]), indent=2))
            return
    print("All current structural sections have a recorded semantic review. Run verify for remaining gates.")


def ocr_page(args):
    project = Path(args.project).resolve()
    with project_lock(project):
        state = load(project)
        e = state["evidence"].get(args.evidence)
        if not e or not e.get("pdf_page"):
            raise ValueError("OCR requires a PDF evidence ID")
        for binary in ("pdftoppm", "tesseract"):
            if not shutil.which(binary):
                raise ValueError(f"{binary} is not installed. Page stays flagged; no OCR was claimed.")
        doc = state["documents"][e["document_id"]]
        if sha(project / doc["source_path"]) != doc["sha256"]:
            raise ValueError("Source copy has changed; OCR refused")
        run = project / EXTRACT / "ocr" / (args.evidence + "-" + dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%S%f"))
        run.mkdir(parents=True)
        subprocess.run(["pdftoppm", "-f", str(e["pdf_page"]), "-l", str(e["pdf_page"]), "-r", "200", "-singlefile", "-png", str(project / doc["source_path"]), str(run / "page")], check=True, timeout=120)
        subprocess.run(["tesseract", str(run / "page.png"), str(run / "ocr"), "txt", "tsv"], check=True, timeout=120)
        state["ocr"][args.evidence] = {"path": rel_path(run, project), "review_status": "unverified", "created_at": now(),
                                      "text_sha256": sha(run / "ocr.txt")}
        # Keep OCR as distinct evidence; never overwrite native extraction or its provenance.
        ocr_id = identifier("EV", [args.evidence, "ocr", sha(run / "ocr.txt")])
        ocr = dict(e, evidence_id=ocr_id, text=(run / "ocr.txt").read_text(encoding="utf-8"),
                   method="tesseract_ocr_unverified", parent_evidence_id=args.evidence,
                   text_path=rel_path(run / "ocr.txt", project), text_file_sha256=sha(run / "ocr.txt"))
        state["evidence"][ocr_id] = ocr
        if ocr_id not in state["sections"][e["section_id"]]["evidence_ids"]:
            state["sections"][e["section_id"]]["evidence_ids"].append(ocr_id)
        state["ocr"][args.evidence]["evidence_id"] = ocr_id
        add_issue(state, "ocr_review", args.evidence, "OCR sidecar generated; inspect words, figures, and numbers against image before use.")
        save(project, state)
        publish(project, state)
        print(json.dumps(state["ocr"][args.evidence]))


def extract_region(project, state, parent_id, bbox, label):
    """Deterministic native crop for columns/answers whose table cell is missing."""
    import pdfplumber
    parent = state["evidence"].get(parent_id)
    if not parent or not parent.get("pdf_page"):
        raise ValueError("Region extraction requires a PDF evidence ID")
    doc = state["documents"][parent["document_id"]]
    if sha(project / doc["source_path"]) != doc["sha256"]:
        raise ValueError("Source integrity failed")
    eid = identifier("EV", [parent_id, "region", bbox])
    if eid in state["evidence"]:
        return eid
    with pdfplumber.open(project / doc["source_path"]) as pdf:
        text = pdf.pages[parent["pdf_page"] - 1].crop(tuple(bbox)).extract_text() or ""
    if not text.strip():
        raise ValueError("Selected region contains no native text; inspect image or OCR")
    path = project / EXTRACT / doc["document_id"] / "regions" / (eid + ".txt")
    write_text(path, text + "\n")
    e = dict(parent, evidence_id=eid, parent_evidence_id=parent_id, text=text, bbox=bbox,
             heading=label, method="pdfplumber_region_unverified", text_path=rel_path(path, project),
             text_file_sha256=sha(path))
    state["evidence"][eid] = e
    state["sections"][parent["section_id"]]["evidence_ids"].append(eid)
    doc["extraction_hashes"][rel_path(path, project)] = sha(path)
    add_issue(state, "region_review", eid, "Verify extracted region bounds, row assignment, and text against the original rendering.")
    return eid


def region_command(args):
    project = Path(args.project).resolve()
    with project_lock(project):
        state = load(project)
        eid = extract_region(project, state, args.evidence, args.bbox, args.label)
        save(project, state)
        publish(project, state)
        print(json.dumps(state["evidence"][eid], indent=2))


def utf8_console():
    """Windows pipes default to a legacy code page; RFP text needs UTF-8."""
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            with contextlib.suppress(ValueError, OSError):
                stream.reconfigure(encoding="utf-8")


def main():
    utf8_console()
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("ingest", help="Copy, hash, extract, index, and detect package changes")
    p.add_argument("--source", required=True)
    p.add_argument("--project", required=True)
    p.add_argument("--name")
    p.add_argument("--exclude", action="append", default=[])
    p = sub.add_parser("import", help="Validate and apply an agent/human semantic review batch")
    p.add_argument("--project", required=True)
    p.add_argument("--batch", required=True)
    for cmd in ("status", "verify", "next"):
        p = sub.add_parser(cmd)
        p.add_argument("--project", required=True)
    p = sub.add_parser("packet", help="Retrieve cited context; estimating mode enforces the live gate")
    p.add_argument("--project", required=True)
    p.add_argument("--query", default="")
    p.add_argument("--mode", choices=["research", "estimating"], default="research")
    p.add_argument("--limit", type=int, default=6)
    p.add_argument("--view", choices=["summary", "full", "blockers"], default="summary")
    p.add_argument("--max-chars", type=int, default=24000, help="Hard serialized character cap, not an exact token limit")
    p.add_argument("--cursor", help="Continue the same query/view/budget; stale cursors are rejected")
    p.add_argument("--record-id", action="append", help="Retrieve specific records plus common constraints and connected context")
    p.add_argument("--evidence-id", action="append", help="Retrieve specific evidence plus connected context")
    p.add_argument("--output")
    p = sub.add_parser("ocr", help="Create an optional separate OCR sidecar for one PDF page")
    p.add_argument("--project", required=True)
    p.add_argument("--evidence", required=True)
    p = sub.add_parser("region", help="Extract a native PDF region with bounding-box provenance")
    p.add_argument("--project", required=True)
    p.add_argument("--evidence", required=True)
    p.add_argument("--bbox", type=float, nargs=4, required=True, metavar=("X0", "TOP", "X1", "BOTTOM"))
    p.add_argument("--label", required=True)
    args = parser.parse_args()
    try:
        if args.command == "ingest":
            ingest(args)
        elif args.command == "import":
            import_batch(args)
        elif args.command == "packet":
            return packet(args)
        elif args.command == "next":
            next_section(args)
        elif args.command == "ocr":
            ocr_page(args)
        elif args.command == "region":
            region_command(args)
        else:
            project = Path(args.project).resolve()
            state = load(project)
            gate = compute_gate(project, state)
            if args.command == "verify":
                with project_lock(project):
                    publish(project, state)
            print(json.dumps(gate, indent=2, ensure_ascii=False))
            return 0 if gate["estimating_allowed"] or args.command == "status" else 2
    except (ValueError, OSError, KeyError, sqlite3.Error, subprocess.SubprocessError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
